#!/usr/bin/env python3
"""CC-EN-GRADING-SHEET Phase A -- the grading workbook.  Run with .corpus-cache/wf/v/bin/python.

    .corpus-cache/wf/v/bin/python tools/en_growth/grading_workbook.py [--seed N] [--out-dir DIR]

Builds ONE .xlsx for a person to grade, and a SEPARATE key file that holds what the
grader must not see (score, Zipf, source group, current tier, carried flag, the
held-out split). Eric chose the separate key over a hidden sheet (2026-10-16):
Excel sheet protection can be undone by anyone, and seeing a score anchors the
grader -- the grades exist to test the score.

The generator takes any row list (build_workbook), so the auditor batch samples and
Eric's spot checks reuse these exact columns (D5). main() builds this file's list:
all low-band words plus a seeded sample of staged words scoring >= 0.62.

Visible columns (identical to what grading_ingest.py reads):
    row_id, word, definition, keep_cut, est_tier, traps, notes
keep_cut is keep / cut / unsure (Eric kept "unsure", 2026-10-16); est_tier is
Easy / Medium / Hard / Expert, exactly as on the earlier auditor sheets. There is no
pronunciation column (Eric's call: ARPAbet is hard to read and he knows the sounds).
"""
import argparse
import csv
import json
import os
import pathlib
import random
import sys

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
from common import ROOT as _ROOT  # noqa: E402
ROOT = pathlib.Path(_ROOT)

KEEP_CUT = ["keep", "cut", "unsure"]
TIERS = ["Easy", "Medium", "Hard", "Expert"]
# Shorthand Eric asked for (2026-10-16): k / c / u, and e / m / h / ex. A tier alone
# means keep. grading_ingest.py expands these; blank in both columns = not graded yet.
KEEP_CUT_SHORT = {"k": "keep", "keep": "keep", "c": "cut", "cut": "cut", "u": "unsure", "unsure": "unsure", "?": "unsure"}
TIER_SHORT = {"e": "Easy", "easy": "Easy", "m": "Medium", "medium": "Medium", "h": "Hard", "hard": "Hard",
              "ex": "Expert", "x": "Expert", "expert": "Expert"}
VISIBLE = ["row_id", "word", "definition", "keep_cut", "est_tier", "traps", "notes"]
HEADER_ROW = 2          # row 1 holds the progress cell
FIRST_DATA_ROW = 3
HELD_OUT = 100
OUT_HOME = pathlib.Path.home() / "repos" / "en-grading"


def order_rows(rows, seed):
    """Seeded shuffle; row_id is the position after shuffling (1-based)."""
    rows = sorted(rows, key=lambda r: r["word"])         # stable input order first
    random.Random(seed).shuffle(rows)
    return [dict(r, row_id=i + 1) for i, r in enumerate(rows)]


def build_workbook(rows, seed, xlsx_path, key_path, title="English grading"):
    """rows: dicts with word, definition, and hidden fields (group, score, zipf,
    current_tier) plus optional carried = {keep_cut, est_tier, traps, notes}."""
    from openpyxl import Workbook
    from openpyxl.formatting.rule import FormulaRule
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.worksheet.datavalidation import DataValidation

    ordered = order_rows(rows, seed)
    n = len(ordered)
    last = FIRST_DATA_ROW + n - 1
    wb = Workbook()
    ws = wb.active
    ws.title = "grade"
    ws["A1"] = "Progress"
    ws["A1"].font = Font(bold=True)
    ws["B1"] = (f'=SUMPRODUCT(--((LEN(D{FIRST_DATA_ROW}:D{last})+LEN(E{FIRST_DATA_ROW}:E{last}))>0))'
                f'&" / {n}"')
    ws["C1"] = (f"{title}. Shortcuts: est_tier e / m / h / ex (a tier alone means keep); "
                "keep_cut c to cut, u if unsure. Ungraded rows are shaded.")
    for c, name in enumerate(VISIBLE, 1):
        cell = ws.cell(row=HEADER_ROW, column=c, value=name)
        cell.font = Font(bold=True)
    for r in ordered:
        car = r.get("carried") or {}
        ws.append([r["row_id"], r["word"], r.get("definition", ""),
                   car.get("keep_cut", ""), car.get("est_tier", ""), car.get("traps", ""), car.get("notes", "")])
    ws.freeze_panes = f"A{FIRST_DATA_ROW}"
    widths = {"A": 8, "B": 20, "C": 70, "D": 11, "E": 11, "F": 30, "G": 30}
    for col, w in widths.items():
        ws.column_dimensions[col].width = w
    for row in ws.iter_rows(min_row=FIRST_DATA_ROW, max_row=last, min_col=3, max_col=3):
        for cell in row:
            cell.alignment = Alignment(wrap_text=True, vertical="top")
    # Dropdowns offer the full values, but typing a shortcut is allowed (no error box).
    dv1 = DataValidation(type="list", formula1='"' + ",".join(KEEP_CUT) + '"', allow_blank=True, showErrorMessage=False)
    dv2 = DataValidation(type="list", formula1='"' + ",".join(TIERS) + '"', allow_blank=True, showErrorMessage=False)
    ws.add_data_validation(dv1)
    ws.add_data_validation(dv2)
    dv1.add(f"D{FIRST_DATA_ROW}:D{last}")
    dv2.add(f"E{FIRST_DATA_ROW}:E{last}")
    ws.conditional_formatting.add(
        f"A{FIRST_DATA_ROW}:G{last}",
        FormulaRule(formula=[f'LEN($D{FIRST_DATA_ROW})+LEN($E{FIRST_DATA_ROW})=0'], fill=PatternFill("solid", fgColor="FFF2CC")))
    # Fixed document metadata, so the same seed gives the same content.
    import datetime
    wb.properties.created = wb.properties.modified = datetime.datetime(2026, 1, 1)
    pathlib.Path(xlsx_path).parent.mkdir(parents=True, exist_ok=True)
    wb.save(xlsx_path)
    rng = random.Random(f"{seed}:held-out")
    held = sorted(rng.sample([r["row_id"] for r in ordered], min(HELD_OUT, n)))
    key = {
        "$doc": "CC-EN-GRADING-SHEET key. Never show this to the grader. Joined to the workbook by row_id.",
        "workbook": pathlib.Path(xlsx_path).name,
        "seed": seed,
        "rows": [{"row_id": r["row_id"], "word": r["word"], "group": r.get("group"), "score": r.get("score"),
                  "zipf": r.get("zipf"), "current_tier": r.get("current_tier"), "carried": bool(r.get("carried"))}
                 for r in ordered],
        "held_out_row_ids": held,
    }
    pathlib.Path(key_path).parent.mkdir(parents=True, exist_ok=True)
    pathlib.Path(key_path).write_text(json.dumps(key, ensure_ascii=False, indent=1) + "\n", encoding="utf-8")
    return ordered, key


def collect_rows(seed, n_staged=300):
    """This file's list: every low-band word plus a seeded staged sample (score >= 0.62).
    Staged words Eric already graded keep their grades, marked carried."""
    from common import zipf_fn
    zipf = zipf_fn()
    home = pathlib.Path(os.path.expanduser("~/repos/spell-wasm"))
    low = list(csv.DictReader(open(ROOT / "reports" / "en-lowband-review-sample.csv", encoding="utf-8")))
    sc = json.load(open(ROOT / "assets" / "words" / "en" / "en_difficulty.json"))["scores"]
    staged = [l.split("\t")[0] for l in (ROOT / "assets" / "words" / "en" / "staged.tsv").read_text(encoding="utf-8").splitlines()[2:] if l.strip()]
    pool = sorted(w for w in staged if sc.get(w, 0) >= 0.62)
    if len(pool) < n_staged:
        raise SystemExit(f"only {len(pool)} staged rows score >= 0.62; not padding (spec non-goal)")
    sample = random.Random(f"{seed}:staged").sample(pool, n_staged)
    gloss = json.load(open(home / ".corpus-cache" / "en-wikt" / "staged_gloss.json", encoding="utf-8"))
    graded = {r["word"]: r for r in csv.DictReader(open(ROOT / "reports" / "en-staged-review-sample-graded.csv", encoding="utf-8"))}
    low_scores = _score_words([r["word"] for r in low])
    rows = []
    for r in low:
        w = r["word"]
        rows.append({"word": w, "definition": r["candidate_1 (flag only, unaudited)"], "group": "low-band",
                     "score": low_scores.get(w), "zipf": round(zipf(w), 2), "current_tier": "not in bank"})
    for w in sample:
        g = (gloss.get(w) or {}).get("glosses") or [""]
        row = {"word": w, "definition": g[0], "group": "staged", "score": sc[w], "zipf": round(zipf(w), 2),
               "current_tier": "staged (untiered)"}
        if w in graded:
            # Eric's 2026-10-16 ruling on that sheet: keep all 200, graded tiers agreed.
            row["carried"] = {"keep_cut": "keep", "est_tier": graded[w]["est_tier"],
                              "traps": graded[w].get("spelling_traps", ""), "notes": graded[w].get("flag", "")}
        rows.append(row)
    return rows


def _score_words(words):
    import subprocess
    import tempfile
    with tempfile.TemporaryDirectory() as d:
        ex = pathlib.Path(d) / "w.txt"
        ex.write_text("\n".join(words) + "\n", encoding="utf-8")
        out = pathlib.Path(d) / "s.json"
        subprocess.run([sys.executable, str(ROOT / "tools" / "en_growth" / "score.py"), "--extra", str(ex), "--out", str(out)],
                       check=True, stdout=subprocess.DEVNULL)
        return json.loads(out.read_text())["scores"]


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--seed", type=int, default=20261016)
    ap.add_argument("--out-dir", default=str(OUT_HOME))
    a = ap.parse_args()
    name = f"en-expert-grading-{a.seed}.xlsx"
    rows = collect_rows(a.seed)
    out = pathlib.Path(a.out_dir)
    ordered, key = build_workbook(rows, a.seed, out / name, out / "keys" / f"{name}.KEY.json")
    carried = sum(1 for r in ordered if r.get("carried"))
    print(f"{out / name}: {len(ordered)} rows ({carried} carried); key {out / 'keys' / (name + '.KEY.json')} (never send it)")


if __name__ == "__main__":
    main()
