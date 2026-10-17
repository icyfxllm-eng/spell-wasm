#!/usr/bin/env python3
"""CC-EN-GRADING-SHEET done items 1-5.  Run: .corpus-cache/wf/v/bin/python tools/en_growth/test_grading.py

Uses a synthetic row list, so it needs no cache or bank inputs.
"""
import contextlib
import io
import json
import os
import pathlib
import sys
import tempfile

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
from grading_workbook import FIRST_DATA_ROW, HEADER_ROW, VISIBLE, build_workbook  # noqa: E402
import grading_ingest  # noqa: E402

fails = []


def check(cond, msg):
    if not cond:
        fails.append(msg)


def fixture_rows(n=20):
    rows = []
    for i in range(n):
        g = "low-band" if i % 2 else "staged"
        rows.append({"word": f"word{i:02d}", "definition": f"meaning {i}", "group": g,
                     "score": 0.6 + i / 100, "zipf": 2.0 + i / 50, "current_tier": "x"})
    return rows


def fill(xlsx, edit):
    """Open the workbook, let edit(ws, row_index, row_id) set cells, save."""
    from openpyxl import load_workbook
    wb = load_workbook(xlsx)
    ws = wb["grade"]
    r = FIRST_DATA_ROW
    while ws.cell(row=r, column=1).value is not None:
        edit(ws, r, ws.cell(row=r, column=1).value)
        r += 1
    wb.save(xlsx)


def run_ingest(args):
    buf = io.StringIO()
    with contextlib.redirect_stdout(buf):
        code = grading_ingest.main(args)
    return code, buf.getvalue()


with tempfile.TemporaryDirectory() as d:
    d = pathlib.Path(d)
    # 1. same seed -> identical order and sample
    a, ka = build_workbook(fixture_rows(), 7, d / "a.xlsx", d / "a.key.json")
    b, kb = build_workbook(fixture_rows(), 7, d / "b.xlsx", d / "b.key.json")
    check([r["word"] for r in a] == [r["word"] for r in b], "1: same seed gave a different row order")
    check(ka["held_out_row_ids"] == kb["held_out_row_ids"], "1: same seed gave a different held-out split")
    c, _ = build_workbook(fixture_rows(), 8, d / "c.xlsx", d / "c.key.json")
    check([r["word"] for r in a] != [r["word"] for r in c], "1: a different seed should change the order")

    # 2. visible sheet: exact columns, exact row count, no hidden data
    from openpyxl import load_workbook
    wb = load_workbook(d / "a.xlsx")
    check(wb.sheetnames == ["grade"], f"2: workbook should hold only the grade sheet, has {wb.sheetnames}")
    ws = wb["grade"]
    head = [c.value for c in ws[HEADER_ROW]]
    check(head[:len(VISIBLE)] == VISIBLE and all(v is None for v in head[len(VISIBLE):]), f"2: columns are {head}")
    data = list(ws.iter_rows(min_row=FIRST_DATA_ROW, values_only=True))
    check(len([r for r in data if r[0] is not None]) == 20, "2: wrong data-row count")
    flat = " ".join(str(v) for row in ws.iter_rows(values_only=True) for v in row if v is not None).lower()
    for leak in ("low-band", "staged", "score", "zipf", "0.61", "current_tier"):
        check(leak not in flat, f"2: hidden data {leak!r} appears on the visible sheet")

    # 3. bad workbook: blank keep_cut, invalid tier, duplicate row_id -> non-zero, all three named
    def bad(ws, r, rid):
        ws.cell(row=r, column=4, value="keep")
        ws.cell(row=r, column=5, value="Hard")
        if rid == 1:
            ws.cell(row=r, column=4).value = None   # cell(value=None) would leave the old value
        if rid == 2:
            ws.cell(row=r, column=5, value="Impossible")
        if rid == 3:
            ws.cell(row=r, column=1, value=4)
    build_workbook(fixture_rows(), 7, d / "bad.xlsx", d / "bad.key.json")
    fill(d / "bad.xlsx", bad)
    code, out = run_ingest(["--workbook", str(d / "bad.xlsx"), "--key", str(d / "bad.key.json"), "--ledger", str(d / "l.jsonl")])
    check(code != 0, "3: a bad workbook must exit non-zero")
    for needle in ("row_id 1 ", "row_id 2 ", "row_id 4: duplicated"):
        check(needle in out, f"3: report does not name {needle!r}:\n{out}")
    check(not (d / "l.jsonl").exists(), "3: a rejected workbook must write nothing")

    # 4. --partial: half graded -> only graded rows written, remainder reported
    def half(ws, r, rid):
        if rid <= 10:
            ws.cell(row=r, column=4, value="keep")
            ws.cell(row=r, column=5, value="Expert" if rid % 3 == 0 else "Medium")
    build_workbook(fixture_rows(), 7, d / "half.xlsx", d / "half.key.json")
    fill(d / "half.xlsx", half)
    code, out = run_ingest(["--workbook", str(d / "half.xlsx"), "--key", str(d / "half.key.json"), "--partial", "--write",
                            "--ledger", str(d / "half.jsonl")])
    led = [json.loads(l) for l in (d / "half.jsonl").read_text().splitlines()] if (d / "half.jsonl").exists() else []
    check(code == 0 and len(led) == 10, f"4: --partial should write 10 rows, wrote {len(led)} (exit {code})")
    check("10 still blank" in out, f"4: should report 10 remaining:\n{out}")

    # 5. fully graded -> three calibration numbers and ledger rows tagged grader=eric
    def full(ws, r, rid):
        ws.cell(row=r, column=4, value="cut" if rid % 5 == 0 else "keep")
        ws.cell(row=r, column=5, value="Expert" if rid % 4 == 0 else "Hard")
    build_workbook(fixture_rows(), 7, d / "full.xlsx", d / "full.key.json")
    fill(d / "full.xlsx", full)
    code, out = run_ingest(["--workbook", str(d / "full.xlsx"), "--key", str(d / "full.key.json"), "--write",
                            "--ledger", str(d / "full.jsonl")])
    check(code == 0, f"5: full workbook should pass:\n{out}")
    for needle in ("Low-band rows graded Expert", "Score precision on the staged sample", "Score band (provisional boundaries)"):
        check(needle in out, f"5: report lacks {needle!r}")
    led = [json.loads(l) for l in (d / "full.jsonl").read_text().splitlines()]
    check(len(led) == 20 and all(e["grader"] == "eric" for e in led), "5: ledger rows must be tagged grader=eric")
    # re-ingesting the same workbook replaces, never duplicates
    run_ingest(["--workbook", str(d / "full.xlsx"), "--key", str(d / "full.key.json"), "--write", "--ledger", str(d / "full.jsonl")])
    check(len((d / "full.jsonl").read_text().splitlines()) == 20, "5: re-ingest duplicated ledger rows")

for f in fails:
    print("FAIL", f)
print(f"test_grading: {'OK' if not fails else str(len(fails)) + ' failure(s)'}")
sys.exit(1 if fails else 0)
