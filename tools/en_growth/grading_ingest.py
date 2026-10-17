#!/usr/bin/env python3
"""CC-EN-GRADING-SHEET Phase B -- read a graded workbook back into the pipeline.

    .corpus-cache/wf/v/bin/python tools/en_growth/grading_ingest.py --workbook X.xlsx --key X.xlsx.KEY.json \\
        [--grader eric] [--partial] [--write] [--ledger PATH] [--report PATH]

Validation (any failure exits non-zero, lists every offending row_id, writes nothing):
  * every row has keep_cut or est_tier (skipped instead with --partial, so grading can span sittings)
  * shorthand accepted: k / c / u for keep_cut; e / m / h / ex for est_tier; a tier alone means keep
  * keep_cut is keep / cut / unsure; every keep row has est_tier in Easy / Medium / Hard / Expert
  * no row_id is missing, duplicated or unknown to the key; the word matches the key

On success with --write, grades go to calibration/en_graded.jsonl (one JSON object per
line), tagged grader, workbook and seed. A row already in the ledger for the same
workbook and row_id is replaced, so re-ingesting a later sitting is safe. "unsure" is
recorded but counts as not graded for calibration. Nothing is promoted or re-tiered here.

Then a calibration report: low-band rows graded Expert, the score's precision on the
staged sample, score band against Eric's tier, and the held-out split from the key.
"""
import argparse
import collections
import json
import os
import pathlib
import sys
import time

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
from common import ROOT as _ROOT  # noqa: E402
ROOT = pathlib.Path(_ROOT)
from grading_workbook import FIRST_DATA_ROW, HEADER_ROW, KEEP_CUT, KEEP_CUT_SHORT, TIERS, TIER_SHORT, VISIBLE  # noqa: E402

LEDGER = ROOT / "calibration" / "en_graded.jsonl"
# Score bands used only to describe the score in the report. These are the
# boundaries the 2026-10-16 refit fit to Eric's first graded sample; the real tier
# boundaries are still unset in config/en_difficulty.json.
PROVISIONAL_BANDS = [(0.513, "Medium"), (0.618, "Hard"), (0.731, "Expert")]


def band(score):
    if score is None:
        return "unscored"
    t = "Easy"
    for cut, name in PROVISIONAL_BANDS:
        if score >= cut:
            t = name
    return t


def read_workbook(path):
    from openpyxl import load_workbook
    ws = load_workbook(path, data_only=True)["grade"]
    head = [c.value for c in ws[HEADER_ROW]][:len(VISIBLE)]
    if head != VISIBLE:
        raise SystemExit(f"workbook columns {head} are not {VISIBLE}")
    rows = []
    for vals in ws.iter_rows(min_row=FIRST_DATA_ROW, max_col=len(VISIBLE), values_only=True):
        if all(v in (None, "") for v in vals):
            continue
        rows.append({k: ("" if v is None else str(v).strip()) for k, v in zip(VISIBLE, vals)})
    return rows


def validate(rows, key, partial):
    """-> (graded rows, problems, blank count)."""
    by_id = {r["row_id"]: r for r in key["rows"]}
    problems, graded, seen, blanks = [], [], collections.Counter(), 0
    for r in rows:
        try:
            rid = int(float(r["row_id"]))
        except ValueError:
            problems.append(f"row_id {r['row_id']!r}: not a number")
            continue
        seen[rid] += 1
        k = by_id.get(rid)
        if k is None:
            problems.append(f"row_id {rid}: unknown to the key")
            continue
        if r["word"] != k["word"]:
            problems.append(f"row_id {rid}: word edited ({r['word']!r}, key says {k['word']!r})")
            continue
        raw_kc, raw_tier = r["keep_cut"].strip().lower(), r["est_tier"].strip().lower()
        kc = KEEP_CUT_SHORT.get(raw_kc, raw_kc)
        tier = TIER_SHORT.get(raw_tier, r["est_tier"].strip())
        if kc == "" and tier in TIERS:
            kc = "keep"   # a tier alone means keep
        if kc == "":
            if partial:
                blanks += 1
                continue
            problems.append(f"row_id {rid} ({k['word']}): not graded (keep_cut and est_tier are both blank)")
            continue
        if kc not in KEEP_CUT:
            problems.append(f"row_id {rid} ({k['word']}): keep_cut {r['keep_cut']!r} is not one of {KEEP_CUT}")
            continue
        if kc == "keep" and tier not in TIERS:
            problems.append(f"row_id {rid} ({k['word']}): est_tier {r['est_tier']!r} is not one of {TIERS}")
            continue
        if tier and tier not in TIERS:
            problems.append(f"row_id {rid} ({k['word']}): est_tier {r['est_tier']!r} is not one of {TIERS}")
            continue
        graded.append({"row_id": rid, "word": k["word"], "keep_cut": kc, "est_tier": tier or None,
                       "traps": r["traps"], "notes": r["notes"], "group": k["group"], "score": k["score"]})
    for rid, c in seen.items():
        if c > 1:
            problems.append(f"row_id {rid}: duplicated ({c} times)")
    missing = sorted(set(by_id) - set(seen))
    for rid in missing:
        problems.append(f"row_id {rid}: missing from the workbook")
    return graded, problems, blanks


def report(graded, key):
    lines = [f"# English grading calibration: {key['workbook']} (seed {key['seed']})", ""]
    usable = [g for g in graded if g["keep_cut"] != "unsure"]
    low = [g for g in usable if g["group"] == "low-band"]
    staged = [g for g in usable if g["group"] == "staged"]
    total_low = sum(1 for r in key["rows"] if r["group"] == "low-band")
    total_staged = sum(1 for r in key["rows"] if r["group"] == "staged")
    low_exp = sum(1 for g in low if g["keep_cut"] == "keep" and g["est_tier"] == "Expert")
    st_exp = sum(1 for g in staged if g["keep_cut"] == "keep" and g["est_tier"] == "Expert")
    lines.append(f"- Low-band rows graded Expert (keep + Expert): **{low_exp}** of {total_low} "
                 f"({len(low)} graded so far)")
    prec = (st_exp / len(staged)) if staged else 0.0
    lines.append(f"- Score precision on the staged sample (share graded Expert): **{prec:.1%}** "
                 f"({st_exp} of {len(staged)} graded; {total_staged} in the sample)")
    lines.append(f"- Unsure rows (not counted): {sum(1 for g in graded if g['keep_cut'] == 'unsure')}")
    lines += ["", "## Score band (provisional boundaries) against your tier", "",
              "| score band \\ your tier | Easy | Medium | Hard | Expert | cut |", "|---|---|---|---|---|---|"]
    conf = collections.Counter()
    for g in usable:
        conf[(band(g["score"]), g["est_tier"] if g["keep_cut"] == "keep" else "cut")] += 1
    for b in ["Easy", "Medium", "Hard", "Expert", "unscored"]:
        cells = [conf[(b, t)] for t in ["Easy", "Medium", "Hard", "Expert", "cut"]]
        if any(cells):
            lines.append(f"| {b} | " + " | ".join(str(c) for c in cells) + " |")
    held = key.get("held_out_row_ids", [])
    lines += ["", f"## Held-out split for the score refit", "",
              f"{len(held)} row_ids, chosen by seed: " + ", ".join(str(h) for h in held)]
    return "\n".join(lines) + "\n"


def main(argv=None):
    ap = argparse.ArgumentParser()
    ap.add_argument("--workbook", required=True)
    ap.add_argument("--key", required=True)
    ap.add_argument("--grader", default="eric")
    ap.add_argument("--partial", action="store_true")
    ap.add_argument("--write", action="store_true")
    ap.add_argument("--ledger", default=str(LEDGER))
    ap.add_argument("--report")
    a = ap.parse_args(argv)
    key = json.loads(pathlib.Path(a.key).read_text(encoding="utf-8"))
    rows = read_workbook(a.workbook)
    graded, problems, blanks = validate(rows, key, a.partial)
    if problems:
        print(f"REJECTED {pathlib.Path(a.workbook).name}: nothing written")
        for p in problems:
            print("  " + p)
        return 1
    print(f"OK {pathlib.Path(a.workbook).name}: {len(graded)} graded"
          + (f", {blanks} still blank" if a.partial else ""))
    rep = report(graded, key)
    print(rep)
    if a.report:
        pathlib.Path(a.report).write_text(rep, encoding="utf-8")
    if not a.write:
        print("dry run: pass --write to record the grades")
        return 0
    ledger = pathlib.Path(a.ledger)
    existing = []
    if ledger.exists():
        existing = [json.loads(l) for l in ledger.read_text(encoding="utf-8").splitlines() if l.strip()]
    wb = pathlib.Path(a.workbook).name
    fresh = {(wb, g["row_id"]) for g in graded}
    kept = [e for e in existing if (e.get("workbook"), e.get("row_id")) not in fresh]
    when = time.strftime("%Y-%m-%d")
    new = [{"word": g["word"], "keep_cut": g["keep_cut"], "est_tier": g["est_tier"], "traps": g["traps"],
            "notes": g["notes"], "grader": a.grader, "workbook": wb, "seed": key["seed"], "row_id": g["row_id"],
            "group": g["group"], "ingested": when} for g in graded]
    ledger.parent.mkdir(parents=True, exist_ok=True)
    ledger.write_text("".join(json.dumps(e, ensure_ascii=False, sort_keys=True) + "\n" for e in kept + new),
                      encoding="utf-8")
    print(f"wrote {len(new)} rows to {ledger}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
